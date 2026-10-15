#!/usr/bin/env python3
"""
=============================================================================
THE LAST DATASET — MASTER ROSTER INTEGRATION SCRIPT
TechTatva 2026 | MAHE Manipal
=============================================================================

Parses the official final registration list:
- 150 teams (1 or 2 members each)
- 222 total participants
- Clean name extraction (strips *(PoC)/Team Leader)
- Deterministic, collision-free TLD-XXXXXX passcodes
- SHA-256 password hashing for Supabase
- Generates:
  1. participants_with_passwords.csv
  2. participants_with_passwords.xlsx
  3. supabase/migrations/006_import_final_150_teams.sql
  4. supabase/APPLY_FINAL_ROSTER.sql
  5. config/teams.json (with 150 allowedIds)
  6. data/progress.json (fallback state for offline testing)
  7. data/teams_roster.json (bundled structured roster)
"""

import os
import sys
import re
import json
import hmac
import hashlib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ALPHABET = 'ABCDEFGHJKLMNPQRSTUVWXYZ23456789'
SALT = b'TLD-TechTatva-Cognitia-2026-MasterRoster'

def get_passcode(team_id: str) -> str:
    h = hmac.new(SALT, team_id.encode('utf-8'), hashlib.sha256).digest()
    code = 'TLD-'
    for i in range(6):
        val = (h[i * 2] << 8) | h[i * 2 + 1]
        code += ALPHABET[val % len(ALPHABET)]
    return code

def hash_passcode(passcode: str) -> str:
    return hashlib.sha256(passcode.strip().encode('utf-8')).hexdigest()

def clean_name(name: str) -> str:
    s = str(name or '').strip()
    s = re.sub(r'\s*\*\s*\(.*?\)\s*/\s*Team\s+Leader\s*', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'\s*\*\s*\(.*?\)', '', s, flags=re.IGNORECASE).strip()
    return s

def main():
    root_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
    
    # Locate source excel file
    candidate_paths = [
        r'c:\Users\Anvesha\OneDrive\Desktop\LAST DATASET\Teams List.xlsx',
        os.path.join(root_dir, 'data', 'Teams List.xlsx'),
        os.path.join(root_dir, 'Teams List.xlsx')
    ]
    
    excel_path = None
    for p in candidate_paths:
        if os.path.exists(p):
            excel_path = p
            break
            
    if not excel_path:
        print(f"Error: Could not locate Teams List.xlsx in candidate paths: {candidate_paths}")
        sys.exit(1)
        
    print(f"Loading workbook from: {excel_path}")
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    
    mahe_keywords = ['manipal', 'international centre for applied sciences', 'icas']
    
    teams = {}
    current_team_id = None
    
    for row in rows[1:]:
        if not any(row):
            continue
        team_id_raw, name, delegate_id, reg_no, college, phone, email, learner_id, category, event = row[:10]
        if team_id_raw and str(team_id_raw).strip():
            current_team_id = str(team_id_raw).strip().upper()
        if not current_team_id:
            continue
            
        if current_team_id not in teams:
            teams[current_team_id] = []
            
        if name or delegate_id:
            raw_name = str(name or '').strip()
            is_poc = 'poc' in raw_name.lower() or 'leader' in raw_name.lower()
            teams[current_team_id].append({
                'raw_name': raw_name,
                'name': clean_name(raw_name),
                'delegate_id': str(delegate_id or '').strip(),
                'reg_no': str(reg_no or '').strip(),
                'college': str(college or '').strip(),
                'phone': str(phone or '').strip(),
                'email': str(email or '').strip(),
                'learner_id': str(learner_id or '').strip(),
                'is_poc': is_poc
            })
            
    print(f"Parsed {len(teams)} distinct teams.")
    
    processed_teams = []
    total_participants = 0
    all_passcodes = set()
    
    for tid, members in teams.items():
        total_participants += len(members)
        passcode = get_passcode(tid)
        all_passcodes.add(passcode)
        passcode_hash = hash_passcode(passcode)
        
        leader_member = next((m for m in members if m['is_poc']), members[0])
        is_mahe = any(any(k in m['college'].lower() for k in mahe_keywords) for m in members)
        student_type = 'MAHE' if is_mahe else 'NONMAHE'
        
        num_suffix = tid.replace('TECHTEAM', '')
        team_name = f"Team {num_suffix}"
        
        formatted_members = []
        for m in members:
            is_leader = (m['delegate_id'] == leader_member['delegate_id'])
            formatted_members.append({
                'delegate_id': m['delegate_id'],
                'name': m['name'],
                'is_leader': is_leader,
                'reg_no': m['reg_no'],
                'college': m['college'],
                'phone': m['phone'],
                'email': m['email'],
                'learner_id': m['learner_id']
            })
            
        processed_teams.append({
            'id': tid,
            'team_name': team_name,
            'team_leader_name': leader_member['name'],
            'leader_mobile': leader_member['phone'],
            'leader_email': leader_member['email'],
            'passcode': passcode,
            'passcode_hash': passcode_hash,
            'student_type': student_type,
            'members': formatted_members
        })
        
    print(f"Total participants processed: {total_participants}")
    print(f"Total unique passcodes: {len(all_passcodes)}")
    assert len(processed_teams) == 150, f"Expected 150 teams, got {len(processed_teams)}"
    assert total_participants == 222, f"Expected 222 participants, got {total_participants}"
    assert len(all_passcodes) == 150, "Passcode collision detected!"
    # All 150 participant teams processed.
    assert len(processed_teams) == 150
    assert total_participants == 222


    # -------------------------------------------------------------------------
    # 1. WRITE STRUCTURED JSON ROSTER
    # -------------------------------------------------------------------------
    json_path = os.path.join(root_dir, 'data', 'teams_roster.json')
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(processed_teams, f, indent=2)
    print(f"Saved: {json_path}")

    # -------------------------------------------------------------------------
    # 2. WRITE ORGANIZER CSV WITH PASSWORDS
    # -------------------------------------------------------------------------
    csv_rows = [
        [
            'Team ID',
            'Team Name',
            'Delegate ID',
            'Member Name',
            'Role',
            'Access Passcode',
            'Mobile No.',
            'Email',
            'College Name',
            'Student Type',
            'Registration No.'
        ]
    ]
    
    for team in processed_teams:
        for m in team['members']:
            csv_rows.append([
                team['id'],
                team['team_name'],
                m['delegate_id'],
                m['name'],
                'Team Leader (PoC)' if m['is_leader'] else 'Member',
                team['passcode'],
                m['phone'],
                m['email'],
                m['college'],
                team['student_type'],
                m['reg_no']
            ])

    def escape_csv(val):
        s = str(val or '')
        if ',' in s or '"' in s or '\n' in s:
            return '"' + s.replace('"', '""') + '"'
        return s

    csv_text = '\n'.join([','.join([escape_csv(col) for col in r]) for r in csv_rows])
    
    # Save in both Desktop/LAST DATASET and project data/
    desktop_folder = r'c:\Users\Anvesha\OneDrive\Desktop\LAST DATASET'
    desktop_csv = os.path.join(desktop_folder, 'participants_with_passwords.csv')
    repo_csv = os.path.join(root_dir, 'data', 'participants_with_passwords.csv')
    
    if os.path.exists(desktop_folder):
        with open(desktop_csv, 'w', encoding='utf-8') as f:
            f.write(csv_text)
        print(f"Saved: {desktop_csv}")
        
    with open(repo_csv, 'w', encoding='utf-8') as f:
        f.write(csv_text)
    print(f"Saved: {repo_csv}")

    # -------------------------------------------------------------------------
    # 3. WRITE FORMATTED EXCEL WORKBOOK
    # -------------------------------------------------------------------------
    out_wb = openpyxl.Workbook()
    ws = out_wb.active
    ws.title = "Official Participants Roster"
    ws.views.sheetView[0].showGridLines = True
    
    headers = [
        'Team ID', 'Team Name', 'Delegate ID', 'Member Name', 'Role', 
        'Access Passcode', 'Mobile No.', 'Email', 'College Name', 'Student Type', 'Registration No.'
    ]
    ws.append(headers)
    
    # Header styling
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28
    
    alt_fill_1 = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    alt_fill_2 = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    leader_badge_font = Font(name="Segoe UI", size=10, bold=True, color="047857")
    passcode_font = Font(name="Consolas", size=11, bold=True, color="B45309")
    regular_font = Font(name="Segoe UI", size=10, color="111827")
    
    current_row = 2
    for team_idx, team in enumerate(processed_teams):
        team_fill = alt_fill_1 if team_idx % 2 == 0 else alt_fill_2
        for m in team['members']:
            ws.append([
                team['id'],
                team['team_name'],
                m['delegate_id'],
                m['name'],
                'Team Leader (PoC)' if m['is_leader'] else 'Member',
                team['passcode'],
                m['phone'],
                m['email'],
                m['college'],
                team['student_type'],
                m['reg_no']
            ])
            
            for c_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=current_row, column=c_idx)
                cell.fill = team_fill
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")
                if c_idx == 1:
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="1E3A8A")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx == 3:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.font = Font(name="Consolas", size=10, color="111827")
                elif c_idx == 5:
                    cell.font = leader_badge_font if m['is_leader'] else regular_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx == 6:
                    cell.font = passcode_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx == 10:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="059669" if team['student_type'] == 'MAHE' else "4F46E5")
                else:
                    cell.font = regular_font
                    
            ws.row_dimensions[current_row].height = 22
            current_row += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            v = str(cell.value or '')
            if len(v) > max_len:
                max_len = len(v)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    desktop_xlsx = os.path.join(desktop_folder, 'participants_with_passwords.xlsx')
    repo_xlsx = os.path.join(root_dir, 'data', 'participants_with_passwords.xlsx')
    
    if os.path.exists(desktop_folder):
        out_wb.save(desktop_xlsx)
        print(f"Saved: {desktop_xlsx}")
    out_wb.save(repo_xlsx)
    print(f"Saved: {repo_xlsx}")

    # -------------------------------------------------------------------------
    # 4. GENERATE PRODUCTION SQL SCRIPT
    # -------------------------------------------------------------------------
    sql_lines = [
        "-- =========================================================================",
        "-- THE LAST DATASET — MASTER 150-TEAM ROSTER IMPORT",
        "-- TechTatva 2026 | MAHE Manipal",
        f"-- Total Teams: {len(processed_teams)} | Total Participants: {total_participants}",
        "-- Generated automatically with SHA-256 passcodes and atomic zone init",
        "-- =========================================================================",
        "",
        "BEGIN;",
        "",
        "-- 0. Purge any non-official test/staging teams and any mock test IDs",
        "DELETE FROM public.teams WHERE id NOT LIKE 'TECHTEAM%' OR id LIKE 'TECHTEAM99%';",
        "DELETE FROM public.score_ledger WHERE team_id NOT LIKE 'TECHTEAM%' OR team_id LIKE 'TECHTEAM99%' OR team_id = 'TECHTEAM10328';",
        "DELETE FROM public.dataset_entitlements WHERE team_id NOT LIKE 'TECHTEAM%' OR team_id LIKE 'TECHTEAM99%' OR team_id = 'TECHTEAM10328';",
        "DELETE FROM public.ml_submissions WHERE team_id NOT LIKE 'TECHTEAM%' OR team_id LIKE 'TECHTEAM99%' OR team_id = 'TECHTEAM10328';",
        "DELETE FROM public.team_sessions WHERE team_id NOT LIKE 'TECHTEAM%' OR team_id LIKE 'TECHTEAM99%' OR team_id = 'TECHTEAM10328';",
        "DELETE FROM public.zone_progress WHERE team_id LIKE 'TECHTEAM99%' OR team_id = 'TECHTEAM10328';",
        ""
    ]
    
    # Teams table inserts
    sql_lines.append("-- 1. Insert / Upsert 150 Teams")
    sql_lines.append("INSERT INTO public.teams (id, team_name, team_leader_name, leader_mobile, passcode_hash, student_type, is_active)")
    sql_lines.append("VALUES")
    
    team_val_lines = []
    for t in processed_teams:
        t_id = t['id'].replace("'", "''")
        t_name = t['team_name'].replace("'", "''")
        t_leader = t['team_leader_name'].replace("'", "''")
        t_mobile = t['leader_mobile'].replace("'", "''")
        t_hash = t['passcode_hash']
        t_type = t['student_type']
        team_val_lines.append(f"('{t_id}', '{t_name}', '{t_leader}', '{t_mobile}', '{t_hash}', '{t_type}', TRUE)")
        
    sql_lines.append(',\n'.join(team_val_lines))
    sql_lines.append("ON CONFLICT (id) DO UPDATE SET")
    sql_lines.append("    team_name = EXCLUDED.team_name,")
    sql_lines.append("    team_leader_name = EXCLUDED.team_leader_name,")
    sql_lines.append("    leader_mobile = EXCLUDED.leader_mobile,")
    sql_lines.append("    passcode_hash = EXCLUDED.passcode_hash,")
    sql_lines.append("    student_type = EXCLUDED.student_type,")
    sql_lines.append("    is_active = TRUE,")
    sql_lines.append("    updated_at = NOW();")
    sql_lines.append("")
    
    # Team members table inserts with trigger hardening
    sql_lines.append("-- 2. Update check_team_member_cap trigger to exclude current delegate_id on upsert")
    sql_lines.append("CREATE OR REPLACE FUNCTION public.check_team_member_cap()")
    sql_lines.append("RETURNS TRIGGER LANGUAGE plpgsql SECURITY DEFINER SET search_path = public, pg_temp AS $$")
    sql_lines.append("DECLARE v_member_count INTEGER;")
    sql_lines.append("BEGIN")
    sql_lines.append("    SELECT COUNT(*) INTO v_member_count FROM public.team_members WHERE team_id = NEW.team_id AND delegate_id <> NEW.delegate_id;")
    sql_lines.append("    IF v_member_count >= 2 THEN")
    sql_lines.append("        RAISE EXCEPTION 'TEAM_CAP_EXCEEDED: Maximum 2 players permitted per team.';")
    sql_lines.append("    END IF;")
    sql_lines.append("    RETURN NEW;")
    sql_lines.append("END; $$;")
    sql_lines.append("")
    sql_lines.append("-- 3. Insert / Upsert 222 Team Members")
    sql_lines.append("INSERT INTO public.team_members (delegate_id, team_id, member_name, is_leader)")
    sql_lines.append("VALUES")
    
    member_val_lines = []
    for t in processed_teams:
        for m in t['members']:
            d_id = m['delegate_id'].replace("'", "''")
            t_id = t['id'].replace("'", "''")
            m_name = m['name'].replace("'", "''")
            is_lead = 'TRUE' if m['is_leader'] else 'FALSE'
            member_val_lines.append(f"('{d_id}', '{t_id}', '{m_name}', {is_lead})")
            
    sql_lines.append(',\n'.join(member_val_lines))
    sql_lines.append("ON CONFLICT (delegate_id) DO UPDATE SET")
    sql_lines.append("    team_id = EXCLUDED.team_id,")
    sql_lines.append("    member_name = EXCLUDED.member_name,")
    sql_lines.append("    is_leader = EXCLUDED.is_leader;")
    sql_lines.append("")
    
    # Zone progress initialization for all 6 zones
    sql_lines.append("-- 4. Initialize Zone Progress for all sectors")
    sql_lines.append("INSERT INTO public.zone_progress (team_id, zone_id, status, unlocked_at)")
    sql_lines.append("SELECT id, 'sanctuary', 'available', NOW() FROM public.teams")
    sql_lines.append("ON CONFLICT (team_id, zone_id) DO NOTHING;")
    sql_lines.append("")
    for locked_zone in ['yellow', 'blue', 'orange', 'red', 'purple']:
        sql_lines.append(f"INSERT INTO public.zone_progress (team_id, zone_id, status, unlocked_at)")
        sql_lines.append(f"SELECT id, '{locked_zone}', 'locked', NULL FROM public.teams")
        sql_lines.append(f"ON CONFLICT (team_id, zone_id) DO NOTHING;")
        sql_lines.append("")

    sql_lines.append("COMMIT;")
    sql_lines.append("")
    
    sql_text = '\n'.join(sql_lines)
    
    sql_file_1 = os.path.join(root_dir, 'supabase', 'migrations', '006_import_final_150_teams.sql')
    sql_file_2 = os.path.join(root_dir, 'supabase', 'APPLY_FINAL_ROSTER.sql')
    
    with open(sql_file_1, 'w', encoding='utf-8') as f:
        f.write(sql_text)
    print(f"Saved: {sql_file_1}")
    
    with open(sql_file_2, 'w', encoding='utf-8') as f:
        f.write(sql_text)
    print(f"Saved: {sql_file_2}")

    # -------------------------------------------------------------------------
    # 5. UPDATE config/teams.json
    # -------------------------------------------------------------------------
    teams_config_path = os.path.join(root_dir, 'config', 'teams.json')
    teams_config = {
        "mode": "open",
        "allowedIds": [t['id'] for t in processed_teams],
        "teamIdPattern": "^[A-Za-z0-9-]{3,32}$",
        "teamNamePattern": "^[\\p{L}\\p{N} .,'’_()-]{2,40}$",
        "notes": f"Configured with {len(processed_teams)} official registered participant teams. IDs are normalized uppercase."
    }
    with open(teams_config_path, 'w', encoding='utf-8') as f:
        json.dump(teams_config, f, indent=2)
    print(f"Saved: {teams_config_path}")

    # -------------------------------------------------------------------------
    # 6. UPDATE data/progress.json (Fallback Store)
    # -------------------------------------------------------------------------
    progress_path = os.path.join(root_dir, 'data', 'progress.json')
    clean_teams = {}
    for team in processed_teams:
        tid = team['id']
        first_member = team['members'][0] if team['members'] else {}
        clean_teams[tid] = {
            "id": tid,
            "name": team["team_name"],
            "passcode": team["passcode"],
            "college": first_member.get('college', ''),
            "createdAt": "2026-10-14T00:00:00.000Z",
            "solvedCtfs": [],
            "events": [],
            "members": [m['delegate_id'] for m in team['members']],
            "memberDetails": team['members'],
            "studentType": team['student_type'],
            "leaderName": team['team_leader_name'],
            "leaderMobile": team['leader_mobile'],
            "leaderEmail": team['leader_email'],
            "completedZones": [],
            "score": 0,
            "sanctuaryCompleted": False
        }

    fallback_data = {
        "teams": clean_teams,
        "sessions": {},
        "submissions": [],
        "scores": {},
        "usedSignatures": [],
        "schemaVersion": 1
    }
        
    with open(progress_path, 'w', encoding='utf-8') as f:
        json.dump(fallback_data, f, indent=2)
    print(f"Saved: {progress_path}")
    
    print("\n=========================================================================")
    print("SUCCESS: 150 TEAMS AND 222 PARTICIPANTS FULLY INTEGRATED & EXPORTED!")
    print("=========================================================================")

if __name__ == '__main__':
    main()
